"""All built-in macro commands.

Every command has the signature fn(rt, args, opts) and returns str, an Action,
a list of Actions, or None.  Flags arrive in opts with value "1".
"""
from __future__ import annotations

import base64
import getpass
import html
import json
import os
import platform
import re
import sqlite3
import urllib.parse
import uuid as _uuid
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from phraseforge.host import (
    Cancelled, CursorAction, FormField, ImageAction, KeyAction, MouseAction, SleepAction, TextAction,
)
from phraseforge.macros import datetimes as dtm
from phraseforge.macros.registry import MacroError, macro, register
from phraseforge.macros.safe_calc import calc as _calc, fmt_number


def _arg(args, i, default=""):
    return args[i] if len(args) > i else default


def _flag(opts, key) -> bool:
    v = opts.get(key)
    return v is not None and v.strip().lower() not in ("0", "false", "no", "off")


def _esc(s: str) -> str:
    """Interpret \\n and \\t in replacement-style arguments."""
    return s.replace("\\n", "\n").replace("\\t", "\t")


def _weekend(rt) -> set[int]:
    return dtm.weekend_days(rt.setting("date.weekend", "sat,sun"))


# ====================================================================== blocks (documentation only)
register("if", "Logic", "{#if cond}...{#elseif cond}...{#else}...{#endif}",
         "Conditional block. Operators: == != < > <= >= contains !contains startswith endswith =~ !~ in, "
         "unary empty/notempty, and/or/not, parentheses. Operators must be unquoted.",
         '{#if {#var name} == "value"}\n\n{#else}\n\n{#endif}')
register("loop", "Logic", "{#loop N -join sep}...{#endloop}",
         "Repeat the body N times. Use {#loopindex} inside.", "{#loop 3}{#loopindex}\n{#endloop}")
register("foreach", "Logic", '{#foreach "a,b,c" -sep , -join sep}...{#endforeach}',
         "Repeat the body for each item; {#item} is the current item.",
         '{#foreach "a,b,c"}{#item}\n{#endforeach}')
register("form", "Forms", '{#form -title "Title"}fields...{#endform}',
         "Show one dialog with all field macros inside the block; read values later with {#formvalue name}.",
         '{#form -title "Details"}\n{#formedit name -label "Name"}\n{#endform}Hello {#formvalue name}')


@macro("loopindex", "Logic", "{#loopindex -zero}", "Current loop iteration (1-based; -zero for 0-based).")
def _loopindex(rt, args, opts):
    if not rt.loops:
        return ""
    i = rt.loops[-1].index
    return str(i if _flag(opts, "zero") else i + 1)


@macro("loopcount", "Logic", "{#loopcount}", "Number of iterations of the current loop.")
def _loopcount(rt, args, opts):
    return str(rt.loops[-1].count) if rt.loops else ""


@macro("item", "Logic", "{#item}", "Current item of {#foreach}.")
def _item(rt, args, opts):
    return rt.loops[-1].item if rt.loops else ""


# ====================================================================== date / time
def _dt_out(rt, dt, opts, default_fmt):
    return dtm.format_dt(dt, opts.get("format", default_fmt), opts.get("lang", rt.setting("date.lang", "en")),
                         opts.get("calendar", rt.setting("date.calendar", "gregorian")),
                         opts.get("digits", rt.setting("date.digits", "latin")))


_WEEKDAYS = {"mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6}


def _shift_weekday(dt, spec: str):
    """-weekday "mon:2" -> the 2nd next Monday (0 = keep date). "mon" = next Monday."""
    from datetime import timedelta
    day, _, n = spec.strip().lower().partition(":")
    if day[:3] not in _WEEKDAYS:
        raise MacroError(f"bad weekday {spec!r}")
    count = int(n or "1")
    if count <= 0:
        return dt
    target = _WEEKDAYS[day[:3]]
    ahead = (target - dt.weekday()) % 7 or 7
    return dt + timedelta(days=ahead + 7 * (count - 1))


def _round_dt(dt, minutes: int, up_only: bool):
    """Round to N minutes (nearest, or up with -roundup)."""
    from datetime import timedelta
    if minutes <= 0:
        return dt
    base = dt.replace(second=0, microsecond=0)
    total = base.hour * 60 + base.minute
    rem = total % minutes
    if rem == 0 and dt.second == 0 and dt.microsecond == 0:
        return base
    if up_only or rem * 2 >= minutes:
        return base + timedelta(minutes=minutes - rem)
    return base - timedelta(minutes=rem)


def _now(rt, opts):
    dt = rt.now()
    if _flag(opts, "utc"):
        dt = datetime.utcnow() if rt.engine.clock is datetime.now else dt
    dt = dtm.apply_offset(dt, opts.get("offset", ""), _weekend(rt))
    if opts.get("weekday"):
        dt = _shift_weekday(dt, opts["weekday"])
    if opts.get("round"):
        dt = _round_dt(dt, int(opts["round"]), _flag(opts, "roundup"))
    return dt


@macro("date", "Date & time", '{#date -format "yyyy-MM-dd" -offset +3d -lang en -calendar hijri -digits arabic}',
       "Current date. Offset units: y, m/mo (month), w, d, wd (workday), h, min, s. "
       "-calendar hijri gives the tabular Hijri date (may differ a day from Umm al-Qura); "
       "-digits arabic writes Arabic-Indic digits.", "{#date -format \"yyyy-MM-dd\"}")
def _date(rt, args, opts):
    return _dt_out(rt, _now(rt, opts), opts, rt.setting("date.format", "yyyy-MM-dd"))


@macro("time", "Date & time", '{#time -format "HH:mm" -offset +30min}', "Current time.", '{#time -format "HH:mm"}')
def _time(rt, args, opts):
    return _dt_out(rt, _now(rt, opts), opts, rt.setting("time.format", "HH:mm"))


@macro("datetime", "Date & time", '{#datetime -format "yyyy-MM-dd HH:mm"}', "Current date and time.")
def _datetime(rt, args, opts):
    return _dt_out(rt, _now(rt, opts), opts, "yyyy-MM-dd HH:mm")


@macro("timestamp", "Date & time", "{#timestamp}", "Unix timestamp (seconds).")
def _timestamp(rt, args, opts):
    return str(int(_now(rt, opts).timestamp()))


@macro("datecalc", "Date & time", '{#datecalc 2026-01-31 -offset +1m -format "dd.MM.yyyy" -diff 2026-03-01 -unit days}',
       "Date arithmetic on a given date (ISO, dd.MM.yyyy or 'today'). With -diff returns the difference "
       "(unit days|weeks|workdays).", "{#datecalc today -offset +7d}")
def _datecalc(rt, args, opts):
    base = dtm.parse_date(_arg(args, 0, "today"), rt.now())
    if "diff" in opts:
        other = dtm.parse_date(opts["diff"], rt.now())
        unit = opts.get("unit", "days").lower()
        if unit.startswith("work"):
            return str(dtm.workdays_between(base.date(), other.date(), _weekend(rt)))
        if unit.startswith("year") or unit.startswith("month"):
            a, b = (base, other) if other >= base else (other, base)
            months = (b.year - a.year) * 12 + (b.month - a.month) - (1 if b.day < a.day else 0)
            n = months // 12 if unit.startswith("year") else months
            return str(n if other >= base else -n)
        if unit.startswith(("hour", "minute", "second")):
            secs = (other - base).total_seconds()
            div = 3600 if unit.startswith("hour") else 60 if unit.startswith("minute") else 1
            return fmt_number(secs / div)
        days = (other.date() - base.date()).days
        return fmt_number(days / 7) if unit.startswith("week") else str(days)
    dt = dtm.apply_offset(base, opts.get("offset", ""), _weekend(rt))
    return _dt_out(rt, dt, opts, rt.setting("date.format", "yyyy-MM-dd"))


@macro("timecalc", "Date & time", '{#timecalc 09:30 -offset +45min -format "HH:mm" -diff 17:00}',
       "Time arithmetic. With -diff returns minutes between the two times.", "{#timecalc now -offset +1h}")
def _timecalc(rt, args, opts):
    base = dtm.parse_time(_arg(args, 0, "now"), rt.now())
    if "diff" in opts:
        other = dtm.parse_time(opts["diff"], rt.now())
        secs = (other - base).total_seconds()
        if opts.get("round"):
            step = int(opts["round"]) * 60
            if step > 0:
                q, r = divmod(abs(secs), step)
                q += 1 if (r and (_flag(opts, "roundup") or r * 2 >= step)) else 0
                secs = q * step * (1 if secs >= 0 else -1)
        unit = opts.get("unit", "minutes").lower()
        if unit == "custom" or (opts.get("format") and "unit" not in opts):
            return _fmt_duration(secs, opts.get("format", "HH:mm"))
        if unit.startswith("h"):
            return fmt_number(secs / 3600)
        if unit.startswith("s"):
            return fmt_number(secs)
        return fmt_number(secs / 60)
    dt = dtm.apply_offset(base, opts.get("offset", ""), _weekend(rt))
    if opts.get("round"):
        dt = _round_dt(dt, int(opts["round"]), _flag(opts, "roundup"))
    return _dt_out(rt, dt, opts, rt.setting("time.format", "HH:mm"))


def _fmt_duration(secs: float, fmt: str) -> str:
    """Format a duration with tokens HH/H (total hours), mm/m, ss/s."""
    sign = "-" if secs < 0 else ""
    s = int(round(abs(secs)))
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    out = re.sub(r"HH|H|mm|m|ss|s", lambda mo: {
        "HH": f"{h:02d}", "H": str(h), "mm": f"{m:02d}", "m": str(m), "ss": f"{sec:02d}", "s": str(sec)}[mo.group(0)], fmt)
    return sign + out


@macro("datepick", "Date & time", '{#datepick "Prompt" -format "yyyy-MM-dd" -name var}',
       "Ask the user for a date with a calendar.")
def _datepick(rt, args, opts):
    iso = rt.host.pick_date(_arg(args, 0, "Date"), rt.now().date().isoformat())
    out = dtm.format_dt(datetime.fromisoformat(iso), opts["format"], opts.get("lang", "en")) if opts.get("format") else iso
    if opts.get("name"):
        rt.form_values[opts["name"]] = out
    return out


# ====================================================================== numbers
@macro("calc", "Numbers", "{#calc 2*(3+4) -decimals 2}",
       "Safe arithmetic: + - * / // % ^ ( ) and round abs min max sqrt floor ceil.", "{#calc 1+1}")
def _calcm(rt, args, opts):
    v = _calc(" ".join(args))
    d = opts.get("decimals")
    out = fmt_number(v, int(d) if d not in (None, "") else None)
    return _format_number_text(out, opts)


def _format_number_text(out: str, opts) -> str:
    """Apply -zeros (min integer digits), -thousands and -decsep to a plain number string."""
    if not any(k in opts for k in ("zeros", "thousands", "decsep")):
        return out
    neg = out.startswith("-")
    body = out[1:] if neg else out
    whole, dot, frac = body.partition(".")
    if opts.get("zeros"):
        whole = whole.zfill(int(opts["zeros"]))
    th = opts.get("thousands", "")
    if th:
        th = " " if th.lower() == "space" else th
        groups = []
        while len(whole) > 3:
            groups.insert(0, whole[-3:])
            whole = whole[:-3]
        whole = th.join([whole] + groups) if whole else th.join(groups)
    dec = opts.get("decsep", ".") or "."
    return ("-" if neg else "") + whole + (dec + frac if dot else "")


@macro("random", "Numbers", "{#random 1 100}", "Random integer between min and max (inclusive).", "{#random 1 100}")
def _random(rt, args, opts):
    lo, hi = int(_arg(args, 0, "1")), int(_arg(args, 1, "100"))
    return str(rt.rng.randint(min(lo, hi), max(lo, hi)))


@macro("randomtext", "Numbers", '{#randomtext "red|green|blue" -sep |}', "Random item from a list.",
       '{#randomtext "a|b|c"}')
def _randomtext(rt, args, opts):
    if len(args) == 1:
        items = args[0].split(opts.get("sep", "|"))
    else:
        items = list(args)
    items = [i for i in items if i != ""]
    return rt.rng.choice(items) if items else ""


@macro("counter", "Numbers", "{#counter invoice -start 1 -step 1 -pad 4}",
       "Persistent counter; returns the next value each time.", "{#counter name}")
def _counter(rt, args, opts):
    name = _arg(args, 0, "default")
    start, step = int(opts.get("start", "1")), int(opts.get("step", "1"))
    cur = rt.host.get_var(name, "counter")
    n = int(cur) + step if cur not in (None, "") else start
    rt.host.set_var(name, str(n), "counter")
    pad = int(opts.get("pad", "0") or 0)
    return str(n).zfill(pad)


@macro("number", "Numbers", '{#number 1234.5 -decimals 2 -thousands ","}', "Format a number.")
def _number(rt, args, opts):
    v = float(_arg(args, 0, "0").replace(",", ""))
    d = int(opts.get("decimals", "0"))
    s = f"{v:,.{d}f}"
    th, dec = opts.get("thousands", ","), opts.get("decimal", ".")
    return s.replace(",", "\x00").replace(".", dec).replace("\x00", th)


# ====================================================================== cursor / keys / mouse
@macro("cursor", "Keys & cursor", "{#cursor}", "Place the caret here after expansion.")
def _cursor(rt, args, opts):
    return CursorAction()


_KEY_ALIASES = {"bksp": "backspace", "bs": "backspace", "del": "delete", "esc": "esc", "escape": "esc",
                "pgup": "pageup", "pgdn": "pagedown", "ins": "insert", "return": "enter", "control": "ctrl",
                "lwin": "win", "rwin": "win", "windows": "win", "space": "space"}


def norm_keys(s: str) -> str:
    parts = [p.strip().lower() for p in s.replace(" ", "").split("+") if p.strip()]
    return "+".join(_KEY_ALIASES.get(p, p) for p in parts)


@macro("key", "Keys & cursor", "{#key ctrl+shift+left -repeat 2}",
       "Press a key or chord. Names: enter tab esc backspace delete up down left right home end pageup "
       "pagedown insert space f1-f12 win; modifiers ctrl alt shift win.", "{#key enter}", aliases=("keys",))
def _key(rt, args, opts):
    rep = int(opts.get("repeat", _arg(args, 1, "1")) or 1)
    return KeyAction(norm_keys(_arg(args, 0, "enter")), repeat=max(1, rep))


def _named_key(keyname):
    def fn(rt, args, opts):
        rep = int(_arg(args, 0, opts.get("repeat", "1")) or 1)
        return KeyAction(keyname, repeat=max(1, rep))
    return fn


for _n, _k in [("enter", "enter"), ("tab", "tab"), ("esc", "esc"), ("bksp", "backspace"), ("del", "delete"),
               ("up", "up"), ("down", "down"), ("left", "left"), ("right", "right"), ("home", "home"),
               ("end", "end"), ("pgup", "pageup"), ("pgdn", "pagedown"), ("ins", "insert")] + \
        [(f"f{i}", f"f{i}") for i in range(1, 13)]:
    register(_n, "Keys & cursor", f"{{#{_n} N}}", f"Press {_k} (optionally N times).", f"{{#{_n}}}",
             fn=_named_key(_k))


@macro("nl", "Text", "{#nl N}", "Insert a line break as text (N times).")
def _nl(rt, args, opts):
    return "\n" * int(_arg(args, 0, "1") or 1)


@macro("tabchar", "Text", "{#tabchar}", "Insert a tab character as text.")
def _tabchar(rt, args, opts):
    return "\t"


@macro("sleep", "Keys & cursor", "{#sleep 500}", "Pause output for N milliseconds.", "{#sleep 500}", aliases=("wait",))
def _sleep(rt, args, opts):
    return SleepAction(int(float(_arg(args, 0, "500"))))


@macro("mouse", "Keys & cursor", "{#mouse click -x 100 -y 200}", "Mouse click|rclick|dblclick|move at x,y.")
def _mouse(rt, args, opts):
    kind = _arg(args, 0, "click").lower()
    if kind not in ("click", "rclick", "dblclick", "move"):
        raise MacroError(f"unknown mouse action {kind}")
    x, y = opts.get("x"), opts.get("y")
    return MouseAction(kind, int(x) if x else None, int(y) if y else None)


# ====================================================================== clipboard / context
@macro("clipboard", "Clipboard", "{#clipboard}", "Current clipboard text.", aliases=("insertclipboard",))
def _clipboard(rt, args, opts):
    return rt.host.get_clipboard() or ""


@macro("clipboardhistory", "Clipboard", "{#clipboardhistory 1}", "Clipboard history entry (0 = newest).")
def _cliphist(rt, args, opts):
    return rt.host.clipboard_history(int(_arg(args, 0, "0"))) or ""


@macro("setclipboard", "Clipboard", "{#setclipboard text}", "Put text on the clipboard (no output).")
def _setclip(rt, args, opts):
    rt.host.set_clipboard(" ".join(args))
    return None


@macro("selection", "Clipboard", "{#selection}", "Text that was selected when the phrase was triggered.")
def _selection(rt, args, opts):
    return rt.ctx.selected_text or ""


@macro("group", "Clipboard", "{#group 1}", "Capture group N of the regex trigger (1 = first group).")
def _group(rt, args, opts):
    n = int(_arg(args, 0, "1"))
    g = rt.ctx.regex_groups
    return (g[n - 1] or "") if 1 <= n <= len(g) else ""


@macro("appname", "System", "{#appname}", "Exe name of the focused application.")
def _appname(rt, args, opts):
    return rt.ctx.app_exe or rt.host.focused_app()[0]


@macro("windowtitle", "System", "{#windowtitle}", "Title of the focused window.")
def _wtitle(rt, args, opts):
    return rt.ctx.window_title or rt.host.focused_app()[1]


@macro("env", "System", "{#env USERPROFILE}", "Environment variable.")
def _env(rt, args, opts):
    return os.environ.get(_arg(args, 0), opts.get("default", ""))


@macro("username", "System", "{#username}", "Windows user name.")
def _username(rt, args, opts):
    return getpass.getuser()


@macro("computername", "System", "{#computername}", "Computer name.")
def _computer(rt, args, opts):
    return platform.node()


@macro("uuid", "System", "{#uuid}", "Random UUID.")
def _uuidm(rt, args, opts):
    return str(_uuid.uuid4())


# ====================================================================== input / forms
def _field(rt, kind, args, opts):
    name = _arg(args, 0, opts.get("name", f"field{len(rt.field_specs) + 1}"))
    label = opts.get("label", name)
    default = opts.get("default", "")
    options = [o.strip() for o in opts.get("options", "").split(opts.get("sep", ","))] if opts.get("options") else []
    if kind == "checkbox":
        default = _flag(opts, "checked") or _flag(opts, "default")
    elif kind == "date" and not default:
        base = dtm.apply_offset(rt.now(), opts.get("shift", ""), _weekend(rt)) if opts.get("shift") else rt.now()
        default = base.date().isoformat()
    elif kind == "date" and default:
        try:
            default = dtm.parse_date(default, rt.now()).date().isoformat()
        except Exception:
            pass
    elif kind in ("combo", "radio") and not default and options:
        default = options[0]
    f = FormField(kind=kind, name=name, label=label, default=default, options=options)
    rt.field_specs[name] = opts
    if rt.collecting is not None:
        rt.collecting.append(f)
        return None
    if kind == "label":
        return None
    from phraseforge.macros.engine import store_form_values
    values = rt.host.show_form(opts.get("title", label), [f])
    store_form_values(rt, [f], values)
    return rt.form_values.get(name, "")


for _kind, _desc in [("edit", "single-line text"), ("memo", "multi-line text"), ("combo", "drop-down (-options a,b)"),
                     ("checkbox", "checkbox (-on text -off text -checked)"), ("radio", "radio group (-options a,b)"),
                     ("date", "date picker (-format)"), ("number", "number"), ("label", "static label")]:
    register(f"form{_kind}", "Forms", f'{{#form{_kind} name -label "Label" -default x}}',
             f"Form field: {_desc}. Inside {{#form}} it adds a field; alone it asks immediately and outputs the value.",
             f'{{#form{_kind} name -label "Label"}}',
             aliases=(("formeditbox",) if _kind == "edit" else ("formcombobox",) if _kind == "combo" else ()),
             fn=(lambda k: lambda rt, a, o: _field(rt, k, a, o))(_kind))


@macro("formvalue", "Forms", "{#formvalue name}", "Value entered in a form field, input or choose (-name).")
def _formvalue(rt, args, opts):
    n = _arg(args, 0)
    spec = rt.field_specs.get(n, {})
    if "enum_items" in spec and n not in rt.form_values:      # {#formenum} inside {#form}
        items = [x for x in spec["enum_items"].split("|") if x]
        chosen = [it for i, it in enumerate(items) if rt.form_values.get(f"{n}__{i}") not in (None, "", "0")]
        sep, last = _esc(spec.get("delim", ", ")), _esc(spec.get("last", " and "))
        return chosen[0] if len(chosen) == 1 else (sep.join(chosen[:-1]) + last + chosen[-1]) if chosen else ""
    if spec.get("phrases") and n in rt.form_values and rt.form_values[n]:   # {#formphrases}
        from phraseforge.macros.commands_extra import _insert_ex
        return _insert_ex(rt, [rt.form_values[n]], {})
    return rt.form_values.get(n, rt.locals.get(n, ""))


@macro("input", "Forms", '{#input "Prompt" -default x -name var -multiline}', "Ask the user for text.",
       '{#input "Prompt"}')
def _input(rt, args, opts):
    prompt = _arg(args, 0, "Input")
    v = rt.host.ask_input(prompt, opts.get("default", ""), multiline=_flag(opts, "multiline"))
    syntax = opts.get("syntax", "").lower()
    tries = 0
    while syntax and syntax != "none" and not _valid_syntax(v, syntax, opts.get("pattern", "")):
        tries += 1
        if tries > 3:
            raise MacroError(f"input does not match {syntax}")
        v = rt.host.ask_input(f"{prompt}\n(Please enter a valid {syntax})", v,
                              multiline=_flag(opts, "multiline"))
    if opts.get("name"):
        rt.form_values[opts["name"]] = v
    return v


SYNTAX_PATTERNS = {
    "number": r"-?\d+([.,]\d+)?",
    "integer": r"-?\d+",
    "email": r"[^@\s]+@[^@\s]+\.[^@\s]+",
    "date": r"\d{1,4}[./-]\d{1,2}[./-]\d{1,4}",
    "time": r"\d{1,2}:\d{2}(:\d{2})?",
    "phone": r"\+?[\d\s()./-]{5,}",
    "url": r"https?://\S+",
    "notempty": r"(?s).*\S.*",
    "letters": r"[^\W\d_]+",
}


def _valid_syntax(v: str, syntax: str, pattern: str = "") -> bool:
    if syntax == "regex":
        return bool(pattern) and re.fullmatch(pattern, v or "") is not None
    pat = SYNTAX_PATTERNS.get(syntax)
    return True if pat is None else re.fullmatch(pat, (v or "").strip()) is not None


@macro("inputmemo", "Forms", '{#inputmemo "Prompt" -default x -name var}', "Ask the user for multi-line text.")
def _inputmemo(rt, args, opts):
    return _input(rt, args, {**opts, "multiline": "1"})


@macro("choose", "Forms", '{#choose "Prompt" -options "a,b,c" -multi -join ", " -name var}',
       "Let the user pick from a list.", '{#choose "Pick one" -options "a,b,c"}')
def _choose(rt, args, opts):
    if "options" in opts:
        items = [x.strip() for x in opts["options"].split(opts.get("sep", ","))]
    else:
        items = list(args[1:])
    picked = rt.host.choose(_arg(args, 0, "Choose"), items, multi=_flag(opts, "multi"))
    v = opts.get("join", ", ").join(picked or [])
    if opts.get("name"):
        rt.form_values[opts["name"]] = v
    return v


@macro("message", "Forms", '{#message "Text" -title "Title"}', "Show a message box (no output).", aliases=("msgbox",))
def _message(rt, args, opts):
    text, title = " ".join(args), opts.get("title", "")
    buttons = MESSAGE_BUTTONS.get(opts.get("buttons", "ok").lower().replace(" ", "").replace("/", ""), ["OK"])
    if buttons == ["OK"]:
        rt.host.message(text, title)
        clicked = "OK"
    else:
        ask = getattr(rt.host, "ask_buttons", None)
        if ask:
            clicked = ask(text, title, buttons)
        else:
            picked = rt.host.choose(f"{title}\n{text}" if title else text, buttons)
            clicked = picked[0] if picked else buttons[-1]
    if opts.get("name"):
        rt.form_values[opts["name"]] = clicked
    # -onyes "Phrase" / -onno ... insert the phrase linked to the clicked button
    target = opts.get("on" + clicked.lower())
    if target:
        return _insert(rt, [target], {})
    if clicked.lower() == "cancel" and _flag(opts, "abortcancel"):
        raise Cancelled("message box cancelled")
    return None


MESSAGE_BUTTONS = {
    "ok": ["OK"], "okcancel": ["OK", "Cancel"], "yesno": ["Yes", "No"],
    "yesnocancel": ["Yes", "No", "Cancel"], "retrycancel": ["Retry", "Cancel"],
    "abortretryignore": ["Abort", "Retry", "Ignore"],
}


# ====================================================================== ask-once prompts
# The value is asked the first time a variable is seen in an expansion and stored in the
# expansion-local scope; later uses of the same variable reuse it without asking.
def _asked(rt, name):
    return name in rt.locals


@macro("ask", "Forms", '{#ask "var" -prompt "Label" -default "x" -multiline}',
       "Ask once per expansion; stores the answer in a local variable and outputs it.",
       '{#ask "name" -prompt "Name"}')
def _ask(rt, args, opts):
    name = _arg(args, 0, "value")
    if not _asked(rt, name):
        rt.locals[name] = rt.host.ask_input(opts.get("prompt", name), opts.get("default", ""),
                                            multiline=_flag(opts, "multiline"))
    return rt.locals[name]


@macro("askchoice", "Forms", '{#askchoice "var" -prompt "Label" -options "a|b|c"}',
       "Ask once for one option from a list (first option is the default); outputs it.",
       '{#askchoice "size" -prompt "Size" -options "S|M|L"}')
def _askchoice(rt, args, opts):
    name = _arg(args, 0, "choice")
    if not _asked(rt, name):
        items = [x.strip() for x in opts.get("options", "").split(opts.get("sep", "|")) if x.strip()]
        picked = rt.host.choose(opts.get("prompt", name), items, multi=False)
        rt.locals[name] = picked[0] if picked else (items[0] if items else "")
    return rt.locals[name]


@macro("askcheck", "Forms", '{#askcheck "var" -prompt "Label" -checked}',
       'Ask once with a checkbox; stores "1" or "0" and outputs nothing. Test it with {#if {#var "var"} == "1"}.',
       '{#askcheck "include" -prompt "Include it?"}')
def _askcheck(rt, args, opts):
    name = _arg(args, 0, "check")
    if not _asked(rt, name):
        f = FormField(kind="checkbox", name=name, label=opts.get("prompt", name), default=_flag(opts, "checked"))
        vals = rt.host.show_form(opts.get("prompt", name), [f]) or {}
        v = vals.get(name, f.default)
        on = v if isinstance(v, bool) else str(v).strip().lower() not in ("", "0", "false", "no", "off")
        rt.locals[name] = "1" if on else "0"
    return None


@macro("askdate", "Forms", '{#askdate "var" -prompt "Label" -format "yyyy-MM-dd"}',
       "Ask once for a date with a calendar; outputs it.", '{#askdate "due" -prompt "Due date"}')
def _askdate(rt, args, opts):
    name = _arg(args, 0, "date")
    if not _asked(rt, name):
        iso = rt.host.pick_date(opts.get("prompt", name), rt.now().date().isoformat())
        if opts.get("format") and iso:
            iso = dtm.format_dt(datetime.fromisoformat(iso), opts["format"], opts.get("lang", "en"))
        rt.locals[name] = iso
    return rt.locals[name]


# ====================================================================== variables
@macro("set", "Variables", "{#set name value -scope local|global}",
       "Set a variable (local to this expansion by default; global persists).", "{#set name value}")
def _set(rt, args, opts):
    rt.set_var(_arg(args, 0), " ".join(args[1:]) if len(args) > 1 else opts.get("value", ""), opts.get("scope", "local"))
    return None


@macro("var", "Variables", "{#var name -scope global -default x}", "Read a variable.", "{#var name}", aliases=("get",))
def _var(rt, args, opts):
    v = rt.get_var(_arg(args, 0), opts.get("scope", ""))
    return v if v != "" else opts.get("default", "")


@macro("unset", "Variables", "{#unset name -scope global}", "Clear a variable.")
def _unset(rt, args, opts):
    scope = opts.get("scope", "local")
    if scope == "local":
        rt.locals.pop(_arg(args, 0), None)
    else:
        rt.host.set_var(_arg(args, 0), "", scope)
    return None


# ====================================================================== phrases
@macro("insert", "Phrases", '{#insert "Phrase description"}', "Insert another phrase (its macros run too).",
       '{#insert "Description"}')
def _insert(rt, args, opts):
    from phraseforge.macros.engine import MAX_INSERT_DEPTH, html_to_text
    from phraseforge.macros.parser import parse
    key = _arg(args, 0)
    if key.lower() in rt.insert_stack or len(rt.insert_stack) >= MAX_INSERT_DEPTH:
        return ""   # cycle or too deep
    found = rt.host.phrase_content(key)
    if found is None and rt.store is not None:
        p = rt.store.find_phrase(key) or (rt.store.get_phrase(int(key)) if key.isdigit() else None)
        found = (p.content, p.content_format) if p else None
    if found is None:
        raise MacroError(f"phrase '{key}' not found")
    content, fmt = found
    if fmt == "html":
        content = html_to_text(content)
    rt.insert_stack.append(key.lower())
    try:
        return rt.engine.eval_nodes(parse(content), rt)
    finally:
        rt.insert_stack.pop()


# ====================================================================== text functions
@macro("upper", "Text", "{#upper text}", "UPPER CASE.")
def _upper(rt, args, opts):
    return " ".join(args).upper()


@macro("lower", "Text", "{#lower text}", "lower case.")
def _lower(rt, args, opts):
    return " ".join(args).lower()


@macro("title", "Text", "{#title text}", "Title Case Each Word.")
def _title(rt, args, opts):
    return re.sub(r"\b(\w)(\w*)", lambda m: m.group(1).upper() + m.group(2).lower(), " ".join(args))


@macro("capitalize", "Text", "{#capitalize text}", "Capitalize the first letter.")
def _cap(rt, args, opts):
    s = " ".join(args)
    return s[:1].upper() + s[1:]


@macro("trim", "Text", "{#trim text}", "Remove surrounding white space.")
def _trim(rt, args, opts):
    return " ".join(args).strip()


@macro("replace", "Text", '{#replace text old new -i}', "Replace text (\\n and \\t allowed in old/new).")
def _replace(rt, args, opts):
    text, old, new = _arg(args, 0), _esc(_arg(args, 1)), _esc(_arg(args, 2))
    if not old:
        return text
    if _flag(opts, "i"):
        return re.sub(re.escape(old), lambda m: new, text, flags=re.I)
    return text.replace(old, new)


@macro("regexreplace", "Text", '{#regexreplace text pattern replacement -i}', "Regex replace (\\1 groups).")
def _rreplace(rt, args, opts):
    return re.sub(_arg(args, 1), _esc(_arg(args, 2)), _arg(args, 0), flags=re.I if _flag(opts, "i") else 0)


@macro("regexget", "Text", '{#regexget text pattern -group 1 -all -sep ", " -i}', "Extract text matching a regex.")
def _rget(rt, args, opts):
    flags = re.I if _flag(opts, "i") else 0
    g = int(opts.get("group", "0"))
    rx = re.compile(_arg(args, 1), flags)
    if _flag(opts, "all"):
        return _esc(opts.get("sep", ", ")).join(m.group(g) or "" for m in rx.finditer(_arg(args, 0)))
    m = rx.search(_arg(args, 0))
    return (m.group(g) or "") if m else ""


@macro("substr", "Text", "{#substr text start length}", "Substring; start is 1-based (negative counts from end).")
def _substr(rt, args, opts):
    s = _arg(args, 0)
    start = int(_arg(args, 1, "1"))
    i = start - 1 if start > 0 else len(s) + start
    ln = _arg(args, 2)
    return s[max(i, 0): (max(i, 0) + int(ln)) if ln else None]


@macro("length", "Text", "{#length text}", "Number of characters.")
def _length(rt, args, opts):
    return str(len(" ".join(args)))


@macro("padleft", "Text", "{#padleft text width -char 0}", "Pad on the left to width.")
def _padleft(rt, args, opts):
    return _arg(args, 0).rjust(int(_arg(args, 1, "0")), (opts.get("char") or " ")[0])


@macro("padright", "Text", "{#padright text width -char .}", "Pad on the right to width.")
def _padright(rt, args, opts):
    return _arg(args, 0).ljust(int(_arg(args, 1, "0")), (opts.get("char") or " ")[0])


@macro("reverse", "Text", "{#reverse text}", "Reverse the text.")
def _reverse(rt, args, opts):
    return " ".join(args)[::-1]


@macro("urlencode", "Text", "{#urlencode text}", "Percent-encode for URLs.")
def _urlenc(rt, args, opts):
    return urllib.parse.quote(" ".join(args), safe="")


@macro("urldecode", "Text", "{#urldecode text}", "Decode percent-encoding.")
def _urldec(rt, args, opts):
    return urllib.parse.unquote(" ".join(args))


@macro("htmlencode", "Text", "{#htmlencode text}", "Escape < > & \" for HTML.")
def _htmlenc(rt, args, opts):
    return html.escape(" ".join(args))


@macro("base64", "Text", "{#base64 text -decode}", "Base64 encode (or -decode).")
def _b64(rt, args, opts):
    s = " ".join(args)
    if _flag(opts, "decode"):
        return base64.b64decode(s).decode("utf-8", errors="replace")
    return base64.b64encode(s.encode()).decode()


# ====================================================================== files & data
def _read_text(path: str, encoding: str = "utf-8") -> str:
    p = Path(os.path.expandvars(path))
    if not p.is_file():
        raise MacroError(f"file not found: {path}")
    return p.read_text(encoding=encoding or "utf-8", errors="replace")


@macro("textfile", "Files & data", '{#textfile "C:\\path\\file.txt" -line 3 -encoding utf-8}',
       "Contents of a text file (or one 1-based line).", aliases=("embed", "embedfile", "embfile", "wordfile"))
def _textfile(rt, args, opts):
    t = _read_text(_arg(args, 0), opts.get("encoding", "utf-8"))
    if opts.get("line"):
        lines = t.splitlines()
        n = int(opts["line"])
        return lines[n - 1] if 1 <= n <= len(lines) else ""
    return t


@macro("image", "Files & data", '{#image "C:\\path\\picture.png"}', "Paste an image file.", aliases=("imgfile",))
def _image(rt, args, opts):
    p = os.path.expandvars(_arg(args, 0))
    if not Path(p).is_file():
        raise MacroError(f"image not found: {p}")
    return ImageAction(p)


@macro("xls", "Files & data", '{#xls "C:\\data.xlsx" -sheet Sheet1 -cell B2}', "Read one Excel cell.")
def _xls(rt, args, opts):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.expandvars(_arg(args, 0)), read_only=True, data_only=True)
    try:
        ws = _sheet(wb, opts.get("sheet", ""))
        if "fromrow" in opts or "fromcol" in opts:
            r1, r2 = int(opts.get("fromrow", "1")), int(opts.get("torow", opts.get("fromrow", "1")))
            c1, c2 = int(opts.get("fromcol", "1")), int(opts.get("tocol", opts.get("fromcol", "1")))
            rows = []
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True):
                rows.append(_esc(opts.get("colsep", ", ")).join(_cell_text(v) for v in row))
            return _esc(opts.get("rowsep", "\\n")).join(rows)
        v = ws[opts.get("cell", "A1")].value
    finally:
        wb.close()
    return _cell_text(v)


def _cell_text(v) -> str:
    return "" if v is None else fmt_number(v) if isinstance(v, (int, float)) else str(v)


def _sheet(wb, spec: str):
    """Sheet by name, or 1-based number."""
    if not spec:
        return wb.worksheets[0]
    if spec.isdigit():
        i = int(spec) - 1
        if not 0 <= i < len(wb.worksheets):
            raise MacroError(f"no sheet {spec}")
        return wb.worksheets[i]
    return wb[spec]


@macro("xlsfind", "Files & data", '{#xlsfind "C:\\data.xlsx" "key" -sheet Sheet1 -searchcol A -returncol C}',
       "Find a row whose search column equals the key and return another column.")
def _xlsfind(rt, args, opts):
    import openpyxl
    from openpyxl.utils import column_index_from_string
    key = _arg(args, 1)
    sc = column_index_from_string(opts.get("searchcol", "A")) - 1
    rc = column_index_from_string(opts.get("returncol", "B")) - 1
    wb = openpyxl.load_workbook(os.path.expandvars(_arg(args, 0)), read_only=True, data_only=True)
    try:
        ws = _sheet(wb, opts.get("sheet", ""))
        for row in ws.iter_rows(values_only=True):
            if len(row) > sc and row[sc] is not None and str(row[sc]).strip().lower() == key.strip().lower():
                v = row[rc] if len(row) > rc else None
                return "" if v is None else fmt_number(v) if isinstance(v, (int, float)) else str(v)
    finally:
        wb.close()
    return opts.get("default", "")


@macro("xml", "Files & data", '{#xml "file.xml or <xml>" -path "./item/name" -attr id -all -sep ", "}',
       "Extract from XML (ElementTree path syntax).")
def _xml(rt, args, opts):
    src = _arg(args, 0)
    root = ET.fromstring(src) if src.lstrip().startswith("<") else ET.parse(os.path.expandvars(src)).getroot()
    path = opts.get("path", ".")
    els = root.findall(path)
    attr = opts.get("attr")
    vals = [(e.get(attr, "") if attr else (e.text or "")).strip() for e in els]
    if _flag(opts, "all"):
        return _esc(opts.get("sep", ", ")).join(vals)
    return vals[0] if vals else ""


@macro("sql", "Files & data", '{#sql "C:\\db.sqlite" "SELECT name FROM t" -sep ", " -rowsep "\\n"}',
       "Query a SQLite database (read-only).")
def _sql(rt, args, opts):
    path = os.path.expandvars(_arg(args, 0))
    if not Path(path).is_file():
        raise MacroError(f"database not found: {path}")
    con = sqlite3.connect(f"file:{urllib.parse.quote(path)}?mode=ro", uri=True)
    try:
        rows = con.execute(_arg(args, 1)).fetchmany(int(opts.get("limit", "1000")))
    finally:
        con.close()
    sep, rsep = _esc(opts.get("sep", ", ")), _esc(opts.get("rowsep", "\\n"))
    return rsep.join(sep.join("" if v is None else str(v) for v in r) for r in rows)


@macro("http", "Internet", '{#http "https://example.com" -method POST -body "x" -header "A: b; C: d" -timeout 30}',
       "HTTP request; outputs the response body.")
def _http(rt, args, opts):
    headers = {}
    for h in opts.get("header", "").split(";"):
        if ":" in h:
            k, v = h.split(":", 1)
            headers[k.strip()] = v.strip()
    return rt.host.http_request(opts.get("method", "GET").upper(), _arg(args, 0), headers or None,
                                opts.get("body"), float(opts.get("timeout", "30")))


_JPATH = re.compile(r"([^.\[\]]+)|\[(\d+)\]")


@macro("json", "Files & data", '{#json text -path "data.items[0].name"}', "Extract a value from JSON text.")
def _json(rt, args, opts):
    obj = json.loads(_arg(args, 0))
    for key, idx in _JPATH.findall(opts.get("path", "")):
        try:
            obj = obj[int(idx)] if idx else obj[key]
        except (KeyError, IndexError, TypeError):
            return opts.get("default", "")
    if isinstance(obj, (dict, list)):
        return json.dumps(obj, ensure_ascii=False)
    if obj is None:
        return ""
    if isinstance(obj, bool):
        return "true" if obj else "false"
    return fmt_number(obj) if isinstance(obj, (int, float)) else str(obj)


# ====================================================================== programs, files, internet
@macro("run", "System", '{#run "notepad.exe" -args "file.txt" -wait -dir "C:\\" -exitcode}',
       "Start a program. -exitcode outputs its exit code (implies -wait).")
def _run(rt, args, opts):
    wait = _flag(opts, "wait") or _flag(opts, "exitcode")
    code = rt.host.run_program(_arg(args, 0), opts.get("args", " ".join(args[1:])), wait=wait,
                               workdir=opts.get("dir", ""))
    return str(code) if _flag(opts, "exitcode") else None


@macro("open", "System", '{#open "C:\\folder or file or https://url"}', "Open a file, folder or URL.",
       aliases=("openfolder", "url"))
def _open(rt, args, opts):
    rt.host.open_target(" ".join(args))
    return None


@macro("mail", "Internet", '{#mail -to a@b.com -subject "Hi" -body "Text" -cc c@d.com}', "Open a new e-mail.")
def _mail(rt, args, opts):
    to = opts.get("to", _arg(args, 0))
    rt.host.open_target(mail_url(to, opts, rt.setting("mail.method", "mailto"),
                                 rt.setting("mail.utf8", "1") != "0"))
    return None


def mail_url(to: str, opts: dict, method: str = "mailto", utf8: bool = True) -> str:
    """mailto: link (or a Gmail compose URL when Settings > Expert Options says Gmail).
    utf8=False encodes with the Windows ANSI code page for old mail clients."""
    fields = {k: opts[k] for k in ("subject", "body", "cc", "bcc") if opts.get(k)}
    if method == "gmail":
        q = {"view": "cm", "fs": "1", "to": to, "su": fields.get("subject", ""), "body": fields.get("body", ""),
             "cc": fields.get("cc", ""), "bcc": fields.get("bcc", "")}
        return "https://mail.google.com/mail/?" + urllib.parse.urlencode(
            {k: v for k, v in q.items() if v}, quote_via=urllib.parse.quote)
    enc = "utf-8" if utf8 else ("mbcs" if os.name == "nt" else "latin-1")
    quote = lambda s, safe="": urllib.parse.quote(s, safe=safe, encoding=enc, errors="replace")  # noqa: E731
    query = "&".join(f"{k}={quote(v)}" for k, v in fields.items())
    return "mailto:" + quote(to, "@,") + ("?" + query if query else "")


@macro("waitfor", "System", '{#waitfor "Window title" -timeout 10 -abort}',
       "Wait until a window whose title contains the text is focused. -abort cancels on timeout.")
def _waitfor(rt, args, opts):
    ok = rt.host.wait_for_window(_arg(args, 0), float(opts.get("timeout", "10")))
    if not ok and _flag(opts, "abort"):
        raise Cancelled(f"window '{_arg(args, 0)}' did not appear")
    return None


# ====================================================================== AI / OCR
@macro("ai", "AI", '{#ai "prompt" -provider openai|anthropic|gpt4all|ollama|openrouter -model name -system "role" '
       '-preset "Summarize"}',
       "Insert an AI model's answer. -preset uses a prompt from Settings > AI text processing > Custom presets "
       "(the macro text becomes the input).", '{#ai "Rewrite politely: {#clipboard}"}', aliases=("openai", "chatgpt"))
def _ai(rt, args, opts):
    prompt = " ".join(args)
    if opts.get("preset"):
        from phraseforge.ai import preset_model, preset_prompt
        ps = {"ai.presets": rt.setting("ai.presets", "[]")}
        instruction = preset_prompt(ps, opts["preset"])
        prompt = f"{instruction}\n\n{prompt}" if prompt else instruction
        if not opts.get("model") and not opts.get("provider"):
            opts = dict(opts, model=preset_model(ps, opts["preset"]))
    return rt.host.ai_complete(prompt, provider=opts.get("provider", ""), model=opts.get("model", ""),
                               system=opts.get("system", "")).strip()


@macro("translate", "AI", '{#translate text -to Arabic -from English -glossary "Name"}',
       "Translate text using the translation engine from Settings > Language Translation "
       "(ignored words and glossaries are applied).", '{#translate {#clipboard} -to Arabic}')
def _translate(rt, args, opts):
    keys = ("translation.ignored_words", "translation.glossaries", "translation.glossary_langs")
    provider = opts.get("provider", "") or rt.setting("translation.provider", "")

    def free():
        from phraseforge.ai import translate_text      # free service, no key needed
        return translate_text({"translation.provider": "free"}, " ".join(args),
                              opts.get("to", "English"), opts.get("from", "")).strip()
    if provider == "free":
        return free()
    try:
        return _translate_with(rt, args, opts, provider, keys)
    except Exception as e:      # noqa: BLE001
        if "key is not set" in str(e):                 # no key for that engine: use the free one
            return free()
        raise


def _translate_with(rt, args, opts, provider, keys):
    from phraseforge.ai import deepl_translate, translation_system
    if provider == "deepl":
        settings = {k: rt.setting(k, "") for k in keys + ("ai.deepl_key",)}
        return deepl_translate(settings, " ".join(args), opts.get("to", "English"), opts.get("from", "")).strip()
    system = translation_system({k: rt.setting(k, "") for k in keys}, opts.get("to", "English"),
                                opts.get("from", ""), opts.get("glossary", ""))
    return rt.host.ai_complete(" ".join(args), provider=provider, model=opts.get("model", ""),
                               system=system).strip()


@macro("ocr", "OCR", "{#ocr -lang Arabic -linebreaks}",
       "Select a screen region and insert its text (Capture2Text). Languages: Arabic, English, Russian.",
       "{#ocr -lang English}")
def _ocr(rt, args, opts):
    lang = opts.get("lang") or rt.setting("ocr.language", "Auto")
    if _flag(opts, "linebreaks"):
        keep = True
    elif _flag(opts, "nolinebreaks"):
        keep = False
    else:   # Settings > Screen OCR > Keep line breaks
        keep = rt.setting("ocr.keep_line_breaks", "1") != "0"
    return rt.host.ocr_region(language=lang, keep_line_breaks=keep)


@macro("ocrfile", "OCR", '{#ocrfile "C:\\image.png" -lang English}', "OCR an image file.")
def _ocrfile(rt, args, opts):
    lang = opts.get("lang") or rt.setting("ocr.language", "Auto")
    return rt.host.ocr_image(os.path.expandvars(_arg(args, 0)), language=lang)


# ensure TextAction import is used for type checkers
_ = TextAction
